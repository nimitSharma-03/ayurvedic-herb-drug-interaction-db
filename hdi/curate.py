"""Phase 4 curation-sheet generation.

Turns the auto-extracted candidate rows in data/processed/candidates.json
into a priority-sorted sheet for human curators, per docs/CURATION_GUIDE.md.
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SEVERITY_ORDER = [
    "contraindication",
    "antagonism",
    "risk_increase",
    "pharmacokinetic",
    "potentiation",
    "general_interaction",
]

# Severities considered clinically actionable enough to warrant a curator's
# attention first. potentiation/general_interaction are left as "standard"
# because they skew generic co-mention language rather than a specific
# pharmacological claim (see docs/CURATION_GUIDE.md for the full rationale).
PRIORITY_SEVERITIES = {"contraindication", "antagonism", "risk_increase", "pharmacokinetic"}

DRUG_CLASSES = {
    "Warfarin": "AC",
    "Heparin": "AC",
    "Aspirin": "AC",
    "Clopidogrel": "AC",
    "Metformin": "AD",
    "Glimepiride": "AD",
    "Insulin": "AD",
    "Pioglitazone": "AD",
    "Digoxin": "CVS",
    "Atenolol": "CVS",
    "Metoprolol": "CVS",
    "Amlodipine": "CVS",
    "Losartan": "CVS",
}

FIELDNAMES = [
    "herb",
    "drug",
    "drug_class",
    "source_pmid",
    "trigger_matched",
    "trigger_group",
    "negated",
    "confidence",
    "priority_tier",
    "evidence_sentence",
    "verdict",
    "notes",
]

VALID_VERDICTS = {"confirmed", "rejected", "unclear"}

# Files a curator's ingest step must never write to, per docs/PROJECT_SCOPE.md.
PROTECTED_PATHS = {"data/processed/candidates.json", "data/raw/raw_abstracts.json"}


def severity_rank(trigger_group):
    return SEVERITY_ORDER.index(trigger_group)


def priority_tier(row):
    return "priority" if (not row["negated"] and row["trigger_group"] in PRIORITY_SEVERITIES) else "standard"


def sort_key(row):
    # (negated, severity_rank): non-negated before negated, then by severity.
    return (row["negated"], severity_rank(row["trigger_group"]))


def balance_by_drug_class(rows):
    """Round-robin interleave rows across drug classes within a tier so a
    single drug class doesn't dominate the top of the block."""
    buckets = {}
    for row in rows:
        buckets.setdefault(row["drug_class"], []).append(row)
    classes = sorted(buckets.keys())
    result = []
    while any(buckets[c] for c in classes):
        for c in classes:
            if buckets[c]:
                result.append(buckets[c].pop(0))
    return result


def build_sheet(candidates):
    rows = []
    for row in candidates:
        row = dict(row)
        row["drug_class"] = DRUG_CLASSES.get(row["drug"], "UNKNOWN")
        row["priority_tier"] = priority_tier(row)
        row.setdefault("verdict", "")
        row.setdefault("notes", "")
        rows.append(row)

    groups = {}
    for row in rows:
        groups.setdefault(sort_key(row), []).append(row)

    ordered = []
    for key in sorted(groups.keys()):
        ordered.extend(balance_by_drug_class(groups[key]))
    return ordered


def write_csv(rows, out_path):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDNAMES)
        writer.writeheader()
        for row in rows:
            writer.writerow({k: row[k] for k in FIELDNAMES})


HEADER_FILL = PatternFill(start_color="FF1F2937", end_color="FF1F2937", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFFFF")
PRIORITY_FILL = PatternFill(start_color="FFFFE699", end_color="FFFFE699", fill_type="solid")
VERDICT_FILLS = {
    "confirmed": PatternFill(start_color="FFC6EFCE", end_color="FFC6EFCE", fill_type="solid"),
    "rejected": PatternFill(start_color="FFFFC7CE", end_color="FFFFC7CE", fill_type="solid"),
    "unclear": PatternFill(start_color="FFFFEB9C", end_color="FFFFEB9C", fill_type="solid"),
}
EVIDENCE_COLUMN_WIDTH = 70


def write_xlsx(rows, out_path):
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "curation_sheet"

    ws.append(FIELDNAMES)
    for row in rows:
        ws.append([row[k] for k in FIELDNAMES])

    n_rows = len(rows) + 1  # +1 for header
    n_cols = len(FIELDNAMES)
    verdict_col = FIELDNAMES.index("verdict") + 1
    priority_col = FIELDNAMES.index("priority_tier") + 1
    evidence_col = FIELDNAMES.index("evidence_sentence") + 1

    # Header: bold white on dark background, frozen, autofilter.
    for col in range(1, n_cols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Column widths: evidence_sentence wide with wrapping, others fit content.
    for col in range(1, n_cols + 1):
        letter = get_column_letter(col)
        if col == evidence_col:
            ws.column_dimensions[letter].width = EVIDENCE_COLUMN_WIDTH
        else:
            max_len = max(
                [len(FIELDNAMES[col - 1])] + [len(str(row[FIELDNAMES[col - 1]])) for row in rows]
            )
            ws.column_dimensions[letter].width = max_len + 2

    for r in range(2, n_rows + 1):
        ws.cell(row=r, column=evidence_col).alignment = Alignment(wrap_text=True, vertical="top")

    # priority_tier highlighting.
    for r in range(2, n_rows + 1):
        if ws.cell(row=r, column=priority_col).value == "priority":
            ws.cell(row=r, column=priority_col).fill = PRIORITY_FILL

    # verdict dropdown restricted to the valid values (blank allowed).
    dv = DataValidation(
        type="list",
        formula1='"' + ",".join(sorted(VALID_VERDICTS)) + '"',
        allow_blank=True,
        showDropDown=False,
    )
    verdict_letter = get_column_letter(verdict_col)
    verdict_range = f"{verdict_letter}2:{verdict_letter}{n_rows}"
    ws.add_data_validation(dv)
    dv.add(verdict_range)

    # Conditional formatting on verdict: confirmed=green, rejected=red, unclear=yellow.
    for verdict, fill in VERDICT_FILLS.items():
        ws.conditional_formatting.add(
            verdict_range,
            CellIsRule(operator="equal", formula=[f'"{verdict}"'], fill=fill),
        )

    ws.sheet_view.showGridLines = False

    wb.save(out_path)


def main(input_path="data/processed/candidates.json", output_path="data/processed/curation_sheet.csv"):
    candidates = json.loads(Path(input_path).read_text(encoding="utf-8"))
    sheet = build_sheet(candidates)
    write_csv(sheet, output_path)
    xlsx_path = Path(output_path).with_suffix(".xlsx")
    write_xlsx(sheet, xlsx_path)
    priority_count = sum(1 for r in sheet if r["priority_tier"] == "priority")
    return len(sheet), priority_count, output_path


def _read_reviewed_sheet(sheet_path):
    sheet_path = Path(sheet_path)
    with sheet_path.open(newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        missing = {"verdict", "notes"} - set(reader.fieldnames or [])
        if missing:
            raise ValueError(
                f"{sheet_path} is missing column(s) {sorted(missing)}; "
                "regenerate it with hdi.curate's `main()` first"
            )
        return list(reader)


def _validate_verdicts(rows):
    errors = []
    for i, row in enumerate(rows, start=2):  # +1 for header, +1 for 1-indexing
        verdict = (row.get("verdict") or "").strip()
        if verdict and verdict not in VALID_VERDICTS:
            errors.append(
                f"line {i}: invalid verdict {verdict!r} for {row.get('herb')}/{row.get('drug')} "
                f"(pmid {row.get('source_pmid')}); must be one of {sorted(VALID_VERDICTS)}"
            )
    if errors:
        raise ValueError("Invalid verdict value(s) in curation sheet:\n" + "\n".join(errors))


def _to_verified_record(row):
    return {
        "herb": row["herb"],
        "drug": row["drug"],
        "drug_class": row["drug_class"],
        "source_pmid": row["source_pmid"],
        "trigger_matched": row["trigger_matched"],
        "trigger_group": row["trigger_group"],
        "negated": row["negated"].strip().lower() == "true",
        "verdict": row["verdict"].strip(),
        "notes": (row.get("notes") or "").strip(),
        "evidence_sentence": row["evidence_sentence"],
    }


def ingest_verdicts(sheet_path="data/processed/curation_sheet.csv", output_path="data/processed/verified_interactions.json"):
    """Read a curator-reviewed curation sheet and write confirmed rows to a
    separate verified-interactions file. Never touches candidates.json or
    raw_abstracts.json (see docs/PROJECT_SCOPE.md) -- this only reads the sheet and
    writes a new output file.
    """
    normalized_output = str(Path(output_path).as_posix())
    if normalized_output in PROTECTED_PATHS:
        raise ValueError(f"refusing to write verified interactions to protected path: {output_path}")

    rows = _read_reviewed_sheet(sheet_path)
    _validate_verdicts(rows)

    verified = [_to_verified_record(r) for r in rows if (r.get("verdict") or "").strip() == "confirmed"]

    out_path = Path(output_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(verified, indent=2, ensure_ascii=False), encoding="utf-8")

    tally = {"confirmed": 0, "rejected": 0, "unclear": 0, "unreviewed": 0}
    for row in rows:
        verdict = (row.get("verdict") or "").strip()
        tally[verdict if verdict in VALID_VERDICTS else "unreviewed"] += 1

    return len(rows), tally, output_path


if __name__ == "__main__":
    import sys

    if len(sys.argv) > 1 and sys.argv[1] == "ingest":
        reviewed, tally, out = ingest_verdicts()
        print(f"read {reviewed} rows {tally}; wrote {tally['confirmed']} verified interactions to {out}")
    else:
        total, priority, out = main()
        print(f"wrote {total} rows ({priority} priority) to {out}")
